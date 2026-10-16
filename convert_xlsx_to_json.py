"""
把 streams.xlsx（直播場次）＋ songs.xlsx（歌曲目錄）＋ song_records.xlsx（演出紀錄）
轉成網站要讀的 streams.json / songs.json
用法：python convert_xlsx_to_json.py
每位主播的三個 xlsx 放在各自的資料夾，轉出來的 json 也放在同一個資料夾：
	uririn/	→ 幽李鈴（index.html）
	chion/	→ 希翁（chion.html）
資料夾裡沒有 streams.xlsx 就跳過
"""
import openpyxl, json, re, datetime, os

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIRS = ['uririn', 'chion']	# 要轉換的資料夾（相對於 repo 根目錄）

# 直播場次的欄位：note 是私人備註，不會輸出到網站；description 是公開的小卡內容（例如推文內文）


def card_text(v):
	r"""小卡內容：把 Excel 裡的各種寫法整理成網頁可以直接顯示的文字
		- 儲存格裡的換行（Alt + Enter）、直接打的 \n（反斜線加 n）都轉成換行
		- \uD83D\uDDA4 這種從程式或 API 複製來的跳脫碼轉回原本的字（例如 🖤）
	"""
	if v in (None, ''):
		return ''
	s = str(v).replace('\r\n', '\n').replace('\\n', '\n')
	s = re.sub(r'\\u([0-9a-fA-F]{4})', lambda m: chr(int(m.group(1), 16)), s)
	s = s.encode('utf-16', 'surrogatepass').decode('utf-16')	# 表情符號是由兩段跳脫碼組成，這裡合併回一個字
	return s.strip()


def header_map(ws):
	header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
	m = {}
	for idx, name in enumerate(header_row):
		if name is None:
			continue
		m[str(name).strip().lower()] = idx
	return m


def normalize_id(v):
	if v in (None, ''):
		return ''
	try:
		f = float(v)
		if f == int(f):
			return str(int(f))
		return str(f)
	except (ValueError, TypeError):
		return str(v).strip()


def get(row, m, *keys, default=None):
	for k in keys:
		if k in m:
			idx = m[k]
			if idx < len(row):
				v = row[idx]
				if v not in (None, ''):
					return v
	return default


def parse_tags(v):
	if v in (None, ''):
		return []
	return [t.strip() for t in re.split(r'[,\uff0c]', str(v)) if t.strip()]


def time_to_seconds(v):
	if v in (None, ''):
		return 0
	if isinstance(v, datetime.timedelta):
		return int(v.total_seconds())
	if isinstance(v, datetime.time):
		return v.hour * 3600 + v.minute * 60 + v.second
	if isinstance(v, (int, float)):
		return int(v * 86400)
	parts = str(v).split(':')
	try:
		parts = [int(p) for p in parts]
	except ValueError:
		return 0
	seconds = 0
	for p in parts:
		seconds = seconds * 60 + p
	return seconds


def seconds_to_display(sec):
	sec = int(sec)
	h, rem = divmod(sec, 3600)
	m, s = divmod(rem, 60)
	if h:
		return f'{h}:{m:02d}:{s:02d}'
	return f'{m}:{s:02d}'


def date_to_str(v):
	if v in (None, ''):
		return ''
	if isinstance(v, (datetime.datetime, datetime.date)):
		return v.strftime('%Y-%m-%d')
	return str(v).strip()


INVISIBLE_CHARS = re.compile(r'[\u200b\u200c\u200d\u200e\u200f\ufeff]')

def clean_name(v):
	if v is None:
		return ''
	s = INVISIBLE_CHARS.sub('', str(v))
	return re.sub(r'\s+', '', s).strip()


def convert_folder(folder):
	"""轉換一個資料夾裡的三個 xlsx，產生同一個資料夾裡的 streams.json、songs.json"""
	folder_path = os.path.join(BASE_DIR, folder)
	label = f'{folder}/' if folder else ''

	# ========== 1. streams.xlsx：一個工作表 = 一個年份 ==========
	wb_streams = openpyxl.load_workbook(os.path.join(folder_path, 'streams.xlsx'), data_only=True)
	streams = {}
	stream_years = []
	for sheet_name in wb_streams.sheetnames:
		ws = wb_streams[sheet_name]
		m = header_map(ws)
		stream_years.append(sheet_name)
		for row in ws.iter_rows(min_row=2, values_only=True):
			sid = get(row, m, 'id')
			if sid in (None, ''):
				continue
			sid = normalize_id(sid)
			streams[sid] = {
				'id': sid,
				'year': sheet_name,
				'date': date_to_str(get(row, m, 'date')),
				'title': str(get(row, m, 'title', default='') or ''),
				'url': str(get(row, m, 'url', default='') or ''),
				'tags': parse_tags(get(row, m, 'tag', 'tags')),
				'guests': parse_tags(get(row, m, 'feat', 'guests', 'guest', '來賓')),
			}
			# 小卡內容：有填 description 才輸出（note 是私人備註，不輸出）
			description = card_text(get(row, m, 'description'))
			if description:
				streams[sid]['description'] = description

	with open(os.path.join(folder_path, 'streams.json'), 'w', encoding='utf-8') as f:
		json.dump({'streams': streams, 'streamYears': stream_years}, f, ensure_ascii=False)

	# ========== 2. songs.xlsx（目錄）：一個工作表 = 一個語言 ==========
	# 目錄的 id 用語言前綴格式（例如 zh1、ja1、en1），本身已經全域唯一，直接用 id 查就好。
	wb_catalog = openpyxl.load_workbook(os.path.join(folder_path, 'songs.xlsx'), data_only=True)
	catalog_by_id = {}     # norm_id -> {artist, language, name}  ← 主要查找鍵
	name_only_lookup = {}  # clean_name -> {artist, language}（id 對不到時的備援，例如舊格式的數字 id）

	for lang in wb_catalog.sheetnames:
		ws = wb_catalog[lang]
		m = header_map(ws)
		for row in ws.iter_rows(min_row=2, values_only=True):
			local_id = get(row, m, 'id')
			name = get(row, m, 'name')
			if local_id in (None, '') or name in (None, ''):
				continue
			artist = str(get(row, m, 'artist', default='') or '')
			norm_id = normalize_id(local_id)
			entry = {'artist': artist, 'language': lang, 'name': str(name).strip()}
			catalog_by_id[norm_id] = entry
			name_only_lookup.setdefault(clean_name(name), entry)

	# ========== 3. song_records.xlsx（演出紀錄）：一個工作表 = 一個年份 ==========
	wb_records = openpyxl.load_workbook(os.path.join(folder_path, 'song_records.xlsx'), data_only=True)
	songs = {lang: [] for lang in wb_catalog.sheetnames}  # 先照目錄的分頁順序建好空清單，順序才會跟 Excel 一致
	unmatched = []

	for sheet_name in wb_records.sheetnames:
		if sheet_name in ('歌曲彙總',):
			continue  # 這個分頁只是 Excel 裡用來做下拉選單的彙總清單，不是演出紀錄本體，跳過
		ws = wb_records[sheet_name]
		m = header_map(ws)
		for row in ws.iter_rows(min_row=2, values_only=True):
			record_id = get(row, m, 'id')
			song_name = get(row, m, 'song_name')
			if record_id in (None, '') or song_name in (None, ''):
				continue

			local_song_id = get(row, m, 'song_id')
			norm_song_id = normalize_id(local_song_id)

			catalog_entry = catalog_by_id.get(norm_song_id)
			if catalog_entry is None:
				# 保險：萬一是還沒改成新格式 id 的舊資料，退回用歌名比對
				catalog_entry = name_only_lookup.get(clean_name(song_name))

			if catalog_entry is None:
				unmatched.append((sheet_name, record_id, song_name))
				artist, language = '', '未分類'
			else:
				artist, language = catalog_entry['artist'], catalog_entry['language']

			time_val = get(row, m, 'time')
			seconds = time_to_seconds(time_val)

			# 歌名一律以目錄裡的乾淨名稱為準，不直接沿用 song_records 自己存的文字
			# （Excel 下拉選單選到的文字可能帶 id 當區分標籤，例如「大笑之歌（zh1）」，不能直接拿來顯示）
			display_name = catalog_entry['name'] if catalog_entry else str(song_name).strip()

			entry = {
				'id': normalize_id(record_id),
				'catalogId': norm_song_id,
				'streamId': normalize_id(get(row, m, 'stream_id')),
				'timeSeconds': seconds,
				'timeDisplay': seconds_to_display(seconds) if seconds else '',
				'name': display_name,
				'artist': artist,
				'feat': str(get(row, m, 'feat', default='') or ''),
				'tags': parse_tags(get(row, m, 'tags', 'tag')),
				'note': str(get(row, m, 'note', default='') or ''),
				'url': str(get(row, m, 'url', default='') or ''),
			}
			songs.setdefault(language, []).append(entry)

	with open(os.path.join(folder_path, 'songs.json'), 'w', encoding='utf-8') as f:
		json.dump({'songs': songs}, f, ensure_ascii=False)

	print(f'轉換完成：{label}streams.json, {label}songs.json')
	if unmatched:
		print(f'⚠️ 有 {len(unmatched)} 筆演出紀錄在目錄裡找不到對應的歌曲（已歸類到「未分類」）：')
		for sheet_name, record_id, song_name in unmatched:
			print(f'  - {sheet_name} 分頁 第 {record_id} 筆：{song_name}')


for folder in DATA_DIRS:
	if not os.path.exists(os.path.join(BASE_DIR, folder, 'streams.xlsx')):
		print(f'略過 {folder}：找不到 streams.xlsx')
		continue
	convert_folder(folder)
